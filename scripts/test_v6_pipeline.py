# -*- coding: utf-8 -*-
"""V6 unit tests: composition charts, segment insight schema, Excel sheets."""
from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))

from lib.constants import (
    EXCEL_PERCENT_FORMAT,
    OVERVIEW_SHEET_NAME,
    RESULT_SHEET_NAME,
    SEGMENT_INSIGHT_SHEET_NAME,
    VOC_TYPES,
)
from lib.excel_io import _module_composition_shares, write_analysis_workbook
from lib.segment_insight_codec import (
    parse_segment_insight_output,
    validate_segment_insight_payload,
)
from lib.segment_stats import (
    build_core_segments,
    build_segment_combinations,
    build_segment_stats_payload,
)
from lib.statistics import aggregate_statistics
from lib.summary_codec import format_summary_text, validate_summary_payload


def _item(item_type: str, dim: str, review_row: int, summary: str = "x") -> dict:
    return {
        "item_type": item_type,
        "dimension": dim,
        "merged_dimension": dim,
        "review_row": review_row,
        "extracted_summary": summary,
    }


def _sample_items() -> list[dict]:
    return [
        _item("消费人群", "幼犬", 2, "幼犬在用"),
        _item("消费人群", "小型犬", 2, "小型犬"),
        _item("产品用途", "消耗精力", 2, "耗精力"),
        _item("消费人群", "幼犬", 3, "幼犬2"),
        _item("使用场景", "日常陪伴", 3, "陪伴"),
        _item("消费人群", "小型犬", 4, "小型"),
        _item("消费人群", "多犬家庭", 4, "两只狗"),
        _item("购买动机", "口碑/推荐", 4, "推荐"),
        _item("消费人群", "老年用户", 5, "老人"),
        _item("消费人群", "学生党", 6, "学生"),
        _item("消费人群", "家庭用户", 7, "家庭"),
        _item("用户满意", "耐用", 2, "耐用"),
        _item("未被满足", "易损坏", 4, "坏了"),
    ]


def _valid_insight(allowed: list[str], *, status: str = "unavailable") -> dict:
    segments = []
    for i, name in enumerate(allowed):
        seg = {
            "segment": name,
            "role": "最终使用者" if i == 0 else "购买决策者",
            "review_findings": [f"{name}评论发现"],
            "ai_profile": [f"{name}联网画像或不可用归纳"],
            "core_needs": [f"{name}需求"],
            "sources": [],
        }
        if status == "ok":
            seg["sources"] = [
                {
                    "source_title": "Example",
                    "source_url": "https://example.com",
                    "finding": f"{name}相关发现",
                }
            ]
        segments.append(seg)
    return {
        "external_research_status": status,
        "segments": segments,
        "product_development": {
            "requirements": [
                {
                    "requirement": "安全结构",
                    "target_segments": allowed[:1] or ["x"],
                    "basis": "评论与角色证据",
                }
            ]
            if allowed
            else [],
            "product_moats": [
                {"moat": "耐久体系", "reason": "高频使用场景需要真正做强"}
            ]
            if allowed
            else [],
        },
    }


def test_v4_eight_dims_unchanged() -> None:
    assert len(VOC_TYPES) == 8
    rows = aggregate_statistics(_sample_items(), total_reviews=10, consolidate=False)
    types = {r["item_type"] for r in rows}
    assert "消费人群" in types and "用户满意" in types


def test_chart_composition_sums_to_one() -> None:
    rows = [
        {"dimension": "儿童", "mention_count": 15},
        {"dimension": "挑食儿童", "mention_count": 3},
    ]
    shares = _module_composition_shares(rows)
    assert abs(sum(shares) - 1.0) < 1e-9
    assert abs(shares[0] - 15 / 18) < 1e-9
    assert abs(shares[1] - 3 / 18) < 1e-9


def test_unique_review_and_coverage() -> None:
    core = build_core_segments(_sample_items(), analyzed_reviews=10, top_n=5)
    by_name = {c["segment"]: c for c in core}
    assert by_name["幼犬"]["review_count"] == 2
    assert by_name["幼犬"]["coverage_rate"] == 20.0


def test_top5_and_under5() -> None:
    core = build_core_segments(_sample_items(), analyzed_reviews=10, top_n=5)
    assert len(core) == 5
    few = [
        _item("消费人群", "A", 2),
        _item("消费人群", "B", 3),
        _item("消费人群", "C", 4),
    ]
    assert len(build_core_segments(few, analyzed_reviews=5, top_n=5)) == 3


def test_no_segments_ok() -> None:
    payload = build_segment_stats_payload(
        [_item("产品用途", "消耗精力", 2)],
        analyzed_reviews=5,
        top_n=5,
    )
    assert payload["segment_count"] == 0
    assert validate_segment_insight_payload(
        _valid_insight([], status="skipped"),
        allowed_segments=[],
    ) is None


def test_combination_unordered_dedup_with_count() -> None:
    core = build_core_segments(_sample_items(), analyzed_reviews=10, top_n=5)
    combos = build_segment_combinations(_sample_items(), core)
    # Unordered sorted label with count
    label = "小型犬 + 幼犬（1条）"
    assert label in (combos.get("幼犬") or [])
    assert label in (combos.get("小型犬") or [])
    flat = "｜".join("｜".join(v) for v in combos.values())
    assert "幼犬 + 小型犬" not in flat  # must not use name-first ordering duplicate
    assert "幼犬 + 学生党" not in flat


def test_ai_rejects_unknown_segment() -> None:
    bad = _valid_insight(["外星用户"], status="unavailable")
    err = validate_segment_insight_payload(bad, allowed_segments=["幼犬"])
    assert err and "不在核心人群" in err


def test_empty_allowed_segments_rejects_invented() -> None:
    invented = _valid_insight(["AI自造人群"], status="unavailable")
    err = validate_segment_insight_payload(invented, allowed_segments=[])
    assert err and "不在核心人群" in err


def test_new_schema_role_requirements_moats() -> None:
    ok = _valid_insight(["幼犬"], status="ok")
    assert validate_segment_insight_payload(ok, allowed_segments=["幼犬"]) is None
    assert ok["segments"][0]["role"] == "最终使用者"
    assert isinstance(ok["product_development"]["requirements"][0], dict)
    assert isinstance(ok["product_development"]["product_moats"][0], dict)

    missing_url = json.loads(json.dumps(ok))
    missing_url["segments"][0]["sources"][0]["source_url"] = ""
    err = validate_segment_insight_payload(missing_url, allowed_segments=["幼犬"])
    assert err and "source_url" in err

    skipped_src = json.loads(json.dumps(ok))
    skipped_src["external_research_status"] = "skipped"
    err2 = validate_segment_insight_payload(skipped_src, allowed_segments=["幼犬"])
    assert err2 and "不得填写 sources" in err2

    text, err3 = parse_segment_insight_output(
        json.dumps(_valid_insight(["幼犬"], status="unavailable"), ensure_ascii=False),
        allowed_segments=["幼犬"],
    )
    assert err3 is None and text is not None


def test_excel_sheets_and_rates() -> None:
    payload = {"sections": [{"title": t, "bullets": [f"{t}要点"]} for t in VOC_TYPES]}
    assert validate_summary_payload(payload) is None
    overview = format_summary_text(payload)

    insight = _valid_insight(["幼犬", "小型犬"], status="unavailable")
    core = build_core_segments(_sample_items(), analyzed_reviews=10, top_n=5)
    combos = build_segment_combinations(_sample_items(), core)
    for row in core:
        row["combination_personas"] = combos.get(row["segment"]) or []

    with tempfile.TemporaryDirectory() as td:
        td_path = Path(td)
        src = td_path / "in.xlsx"
        out = td_path / "out.xlsx"
        wb = Workbook()
        wb.active.title = "Reviews"
        wb["Reviews"].append(["标题", "内容"])
        wb["Reviews"].append(["a", "b"])
        wb.save(src)
        wb.close()

        summary_rows = aggregate_statistics(_sample_items(), total_reviews=10)
        # Force two 消费人群 rows for composition check via export meta
        meta = write_analysis_workbook(
            src,
            out,
            summary_rows=summary_rows,
            total_reviews=10,
            overview_summary=overview,
            core_segments=core,
            segment_insight=insight,
        )
        assert "AI总结" not in meta["sheetnames"]
        assert SEGMENT_INSIGHT_SHEET_NAME in meta["sheetnames"]
        assert OVERVIEW_SHEET_NAME in meta["sheetnames"]
        assert RESULT_SHEET_NAME in meta["sheetnames"]
        assert "产品决策分析" not in meta["sheetnames"]

        for m in meta.get("module_metas") or []:
            if m.get("has_chart"):
                assert abs(float(m.get("composition_sum") or 0) - 1.0) < 1e-6

        wb2 = load_workbook(out)
        assert "AI总结" not in wb2.sheetnames
        ov = wb2[OVERVIEW_SHEET_NAME]
        assert ov.cell(row=1, column=15).value == "AI总结"
        assert "消费人群" in str(ov.cell(row=2, column=15).value or "")

        result = wb2[RESULT_SHEET_NAME]
        assert result.cell(row=2, column=4).number_format == EXCEL_PERCENT_FORMAT
        # coverage: unique/analyzed — puppy 2/10 = 0.2 appears somewhere for 幼犬
        rates = [
            float(result.cell(row=r, column=4).value or 0)
            for r in range(2, result.max_row + 1)
            if result.cell(row=r, column=2).value == "幼犬"
        ]
        assert rates and abs(rates[0] - 0.2) < 1e-9

        seg = wb2[SEGMENT_INSIGHT_SHEET_NAME]
        text = "\n".join(
            str(seg.cell(row=r, column=1).value or "")
            for r in range(1, seg.max_row + 1)
        )
        assert "01 核心消费人群" in text
        assert "02 核心人群画像" in text
        assert "03 产品开发方向" in text
        assert "行为特征" not in text
        # role column present
        assert seg.cell(row=2, column=2).value == "角色"
        wb2.close()


def run_all() -> None:
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    for t in tests:
        t()
        print(f"  ok: {t.__name__}")
    print(f"V6_PIPELINE_TESTS_PASSED ({len(tests)})")


if __name__ == "__main__":
    run_all()

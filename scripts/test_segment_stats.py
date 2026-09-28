# -*- coding: utf-8 -*-
"""Tests for simplified consumer segment deep analysis."""
from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook, load_workbook

from lib.constants import (
    AI_SUMMARY_SHEET_NAME,
    DECISION_SHEET_NAME,
    INTELLIGENCE_SCHEMA_VERSION,
    OVERVIEW_SHEET_NAME,
    RESULT_SHEET_NAME,
    SEGMENT_SHEET_NAME,
    VOC_TYPES,
)
from lib.excel_io import write_analysis_workbook
from lib.intelligence_codec import (
    compute_intelligence_input_hash,
    parse_intelligence_output,
    validate_intelligence_payload,
)
from lib.segment_stats import (
    build_segment_analysis,
    discover_segments,
    segment_dim_cooccur,
)


def _items_two_segments() -> list[dict]:
    items = []
    for rid in range(1, 13):
        items.append(
            {
                "item_type": "消费人群",
                "dimension": "大型犬主人",
                "merged_dimension": "大型犬主人",
                "extracted_summary": f"large-{rid}",
                "review_row": rid,
            }
        )
        items.append(
            {
                "item_type": "产品用途",
                "dimension": "高强度撕咬",
                "merged_dimension": "高强度撕咬",
                "extracted_summary": f"chew-{rid}",
                "review_row": rid,
            }
        )
    for rid in range(13, 21):
        items.append(
            {
                "item_type": "消费人群",
                "dimension": "小型犬主人",
                "merged_dimension": "小型犬主人",
                "extracted_summary": f"small-{rid}",
                "review_row": rid,
            }
        )
        items.append(
            {
                "item_type": "产品用途",
                "dimension": "日常陪伴",
                "merged_dimension": "日常陪伴",
                "extracted_summary": f"compan-{rid}",
                "review_row": rid,
            }
        )
    for rid in (1, 2, 3, 4, 5, 6):
        items.append(
            {
                "item_type": "未被满足",
                "dimension": "不耐咬",
                "merged_dimension": "不耐咬",
                "extracted_summary": f"pain-{rid}",
                "review_row": rid,
                "signal_type": "明确问题",
            }
        )
    # Same review hits segment twice — dedup
    items.append(
        {
            "item_type": "消费人群",
            "dimension": "大型犬主人",
            "merged_dimension": "大型犬主人",
            "extracted_summary": "dup",
            "review_row": 1,
        }
    )
    return items


def test_n1_segment_analyzable() -> None:
    items = [
        {
            "item_type": "消费人群",
            "merged_dimension": "通勤者",
            "dimension": "通勤者",
            "extracted_summary": "commute",
            "review_row": 1,
        },
        {
            "item_type": "用户满意",
            "merged_dimension": "便携",
            "dimension": "便携",
            "extracted_summary": "portable",
            "review_row": 1,
            "signal_type": "满意点",
        },
    ]
    analysis = build_segment_analysis(items, analyzed_reviews=10)
    assert analysis["known_segments"] == ["通勤者"]
    seg = analysis["segments"][0]
    assert seg["mention_count"] == 1
    assert seg["mention_rate"] == 10.0
    assert seg["associations"]["用户满意"][0]["dimension"] == "便携"
    assert seg["associations"]["用户满意"][0]["mention_rate"] == 100.0


def test_n_lt_10_can_emit_opportunity() -> None:
    """n<10 must NOT block product opportunity output."""
    items = []
    for rid in range(1, 4):
        items.append(
            {
                "item_type": "消费人群",
                "merged_dimension": "新手",
                "dimension": "新手",
                "extracted_summary": "n",
                "review_row": rid,
            }
        )
        items.append(
            {
                "item_type": "未被满足",
                "merged_dimension": "说明书不足",
                "dimension": "说明书不足",
                "extracted_summary": "doc",
                "review_row": rid,
                "signal_type": "明确需求",
            }
        )
    summary = [
        {"item_type": t, "dimension": "维", "mention_count": 1, "mention_rate": 5.0}
        for t in VOC_TYPES
    ]
    summary[0] = {"item_type": "消费人群", "dimension": "新手", "mention_count": 3, "mention_rate": 30.0}
    summary[7] = {"item_type": "未被满足", "dimension": "说明书不足", "mention_count": 3, "mention_rate": 30.0}
    payload = {
        "sections": [{"title": t, "bullets": ["ok"]} for t in VOC_TYPES],
        "attribute_performance": [],
        "pain_priorities": [],
        "opportunities": [],
        "recommendations": {
            "priority_improvements": [],
            "keep_strengths": [],
            "explore_opportunities": [],
        },
        "segment_intelligence": {
            "segments": [
                {
                    "segment": "新手",
                    "summary": "n=3 仍可解读",
                    "important_attributes": [],
                    "development_implications": ["简化引导"],
                }
            ],
            "comparisons": [],
            "external_research_status": "unavailable",
            "external_research": [],
            "segment_product_opportunities": [
                {
                    "segment": "新手",
                    "opportunity": "加强说明书",
                    "review_evidence": ["说明书不足"],
                    "recommendation": "出简明卡片",
                }
            ],
        },
    }
    seg = build_segment_analysis(items, 10)
    assert seg["segments"][0]["mention_count"] == 3
    enriched, text, err = parse_intelligence_output(
        json.dumps(payload, ensure_ascii=False),
        summary_rows=summary,
        items=items,
        analyzed_reviews=10,
        segment_analysis=seg,
    )
    assert err is None, err
    assert enriched["segment_intelligence"]["segment_product_opportunities"]
    assert "消费人群深度分析" in text


def test_unique_review_dedup() -> None:
    items = _items_two_segments()
    segs = discover_segments(items, 30)
    large = next(s for s in segs if s["segment"] == "大型犬主人")
    assert large["mention_count"] == 12


def test_within_segment_rate_and_share() -> None:
    items = _items_two_segments()
    analyzed = 30
    analysis = build_segment_analysis(items, analyzed)
    by = {s["segment"]: s for s in analysis["segments"]}
    assert by["大型犬主人"]["mention_rate"] == 40.0
    chew = by["大型犬主人"]["associations"]["产品用途"][0]
    assert chew["dimension"] == "高强度撕咬"
    assert chew["mention_rate"] == 100.0  # 12/12


def test_pp_diff() -> None:
    items = _items_two_segments()
    analysis = build_segment_analysis(items, 30)
    chew = [
        c
        for c in analysis["comparisons"]
        if c["dimension"] == "高强度撕咬"
        and {c["segment_a"], c["segment_b"]} == {"大型犬主人", "小型犬主人"}
    ]
    assert chew
    c = chew[0]
    assert abs(c["pp_diff"]) == 100.0
    assert "p_value" not in c
    assert "test" not in c
    assert "sample_status" not in c


def test_opportunity_requires_segment_cooccurrence() -> None:
    # 不耐咬 only on large-dog reviews; assigning to 小型犬主人 must fail
    items = _items_two_segments()
    summary = [
        {"item_type": t, "dimension": "维", "mention_count": 1, "mention_rate": 5.0}
        for t in VOC_TYPES
    ]
    summary[0] = {"item_type": "消费人群", "dimension": "小型犬主人", "mention_count": 8, "mention_rate": 26.7}
    summary.append(
        {"item_type": "消费人群", "dimension": "大型犬主人", "mention_count": 12, "mention_rate": 40.0}
    )
    summary[7] = {"item_type": "未被满足", "dimension": "不耐咬", "mention_count": 6, "mention_rate": 20.0}
    # Fix known dims: both segments present
    known_dims = {t: {"维"} for t in VOC_TYPES}
    known_dims["消费人群"] = {"大型犬主人", "小型犬主人"}
    known_dims["未被满足"] = {"不耐咬", "维"}
    payload = {
        "sections": [{"title": t, "bullets": ["ok"]} for t in VOC_TYPES],
        "attribute_performance": [],
        "pain_priorities": [],
        "opportunities": [],
        "recommendations": {
            "priority_improvements": [],
            "keep_strengths": [],
            "explore_opportunities": [],
        },
        "segment_intelligence": {
            "segments": [],
            "comparisons": [],
            "external_research_status": "skipped",
            "external_research": [],
            "segment_product_opportunities": [
                {
                    "segment": "小型犬主人",
                    "opportunity": "误归痛点",
                    "review_evidence": ["不耐咬"],
                    "recommendation": "不应通过",
                }
            ],
        },
    }
    err = validate_intelligence_payload(
        payload,
        known_dims=known_dims,
        known_segments={"大型犬主人", "小型犬主人"},
        items=items,
    )
    assert err and "共现" in err
    # Correct assignment to 大型犬主人 passes
    assert segment_dim_cooccur(items, "大型犬主人", "不耐咬") >= 1
    assert segment_dim_cooccur(items, "小型犬主人", "不耐咬") == 0
    payload["segment_intelligence"]["segment_product_opportunities"][0]["segment"] = "大型犬主人"
    err2 = validate_intelligence_payload(
        payload,
        known_dims=known_dims,
        known_segments={"大型犬主人", "小型犬主人"},
        items=items,
    )
    assert err2 is None, err2


def test_external_research_unavailable_ok() -> None:
    items = [
        {
            "item_type": "消费人群",
            "merged_dimension": "新手",
            "dimension": "新手",
            "review_row": 1,
            "extracted_summary": "x",
        },
        {
            "item_type": "未被满足",
            "merged_dimension": "说明书不足",
            "dimension": "说明书不足",
            "review_row": 1,
            "extracted_summary": "y",
            "signal_type": "明确需求",
        },
    ]
    summary = [
        {"item_type": t, "dimension": "维", "mention_count": 1, "mention_rate": 5.0}
        for t in VOC_TYPES
    ]
    summary[0]["dimension"] = "新手"
    summary[7]["dimension"] = "说明书不足"
    payload = {
        "sections": [{"title": t, "bullets": ["ok"]} for t in VOC_TYPES],
        "attribute_performance": [],
        "pain_priorities": [],
        "opportunities": [],
        "recommendations": {
            "priority_improvements": [],
            "keep_strengths": [],
            "explore_opportunities": [],
        },
        "segment_intelligence": {
            "segments": [
                {
                    "segment": "新手",
                    "summary": "ok",
                    "important_attributes": [],
                    "development_implications": ["观察"],
                }
            ],
            "comparisons": [],
            "external_research_status": "unavailable",
            "external_research": [],
            "segment_product_opportunities": [
                {
                    "segment": "新手",
                    "opportunity": "简化上手",
                    "review_evidence": ["说明书不足"],
                    "recommendation": "验证引导",
                }
            ],
        },
    }
    seg = build_segment_analysis(items, 10)
    enriched, text, err = parse_intelligence_output(
        json.dumps(payload, ensure_ascii=False),
        summary_rows=summary,
        items=items,
        analyzed_reviews=10,
        segment_analysis=seg,
    )
    assert err is None, err
    assert enriched["segment_intelligence"]["external_research_status"] == "unavailable"


def test_input_hash_follows_prompts() -> None:
    h1 = compute_intelligence_input_hash(
        system_prompt="SYS-A",
        user_prompt="USER-A",
        schema_version=INTELLIGENCE_SCHEMA_VERSION,
    )
    h2 = compute_intelligence_input_hash(
        system_prompt="SYS-B",
        user_prompt="USER-A",
        schema_version=INTELLIGENCE_SCHEMA_VERSION,
    )
    h3 = compute_intelligence_input_hash(
        system_prompt="SYS-A",
        user_prompt="USER-B",
        schema_version=INTELLIGENCE_SCHEMA_VERSION,
    )
    h1b = compute_intelligence_input_hash(
        system_prompt="SYS-A",
        user_prompt="USER-A",
        schema_version=INTELLIGENCE_SCHEMA_VERSION,
    )
    assert h1 == h1b
    assert h1 != h2
    assert h1 != h3


def test_excel_segment_sheet() -> None:
    items = _items_two_segments()
    analysis = build_segment_analysis(items, 30)
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "in.xlsx"
        out = Path(td) / "out.xlsx"
        wb = Workbook()
        wb.active.title = "Reviews"
        wb["Reviews"].append(["标题", "内容"])
        wb.save(src)
        wb.close()
        text = "AI评论洞察总结\n\n【消费人群深度分析】\n· 主要人群：\n  - 大型犬主人（n=12）"
        meta = write_analysis_workbook(
            src,
            out,
            summary_rows=[
                {
                    "item_type": "消费人群",
                    "dimension": "大型犬主人",
                    "mention_count": 12,
                    "mention_rate": 40.0,
                    "representative_feedback": "x",
                    "theme_summary": "「大型犬主人」提及率 40.0%。",
                }
            ],
            total_reviews=30,
            analyzed_reviews=30,
            overview_summary=text,
            intelligence={
                "segment_intelligence": {
                    "segments": [],
                    "comparisons": [],
                    "external_research_status": "unavailable",
                    "external_research": [],
                    "segment_product_opportunities": [],
                }
            },
            segment_analysis=analysis,
        )
        assert meta["sheetnames"][-5:] == [
            OVERVIEW_SHEET_NAME,
            DECISION_SHEET_NAME,
            SEGMENT_SHEET_NAME,
            AI_SUMMARY_SHEET_NAME,
            RESULT_SHEET_NAME,
        ]
        wb2 = load_workbook(out)
        ws = wb2[SEGMENT_SHEET_NAME]
        assert ws["A1"].value == "消费人群深度分析"
        # Overview headers: 人群 / 评论数 n / 占有效评论比例 — no 样本状态
        joined = " ".join(
            str(ws.cell(row=r, column=c).value or "")
            for r in range(1, 40)
            for c in range(1, 9)
        )
        assert "样本状态" not in joined
        assert "Fisher" not in joined and "χ²" not in joined and "p=" not in joined
        assert "消费人群深度分析" in str(wb2[AI_SUMMARY_SHEET_NAME]["A2"].value)
        wb2.close()


def test_reject_unknown_segment() -> None:
    payload = {
        "sections": [{"title": t, "bullets": ["ok"]} for t in VOC_TYPES],
        "attribute_performance": [],
        "pain_priorities": [],
        "opportunities": [],
        "recommendations": {
            "priority_improvements": [],
            "keep_strengths": [],
            "explore_opportunities": [],
        },
        "segment_intelligence": {
            "segments": [
                {
                    "segment": "火星用户",
                    "summary": "x",
                    "important_attributes": [],
                    "development_implications": ["y"],
                }
            ],
            "comparisons": [],
            "external_research_status": "skipped",
            "external_research": [],
            "segment_product_opportunities": [],
        },
    }
    err = validate_intelligence_payload(
        payload,
        known_dims={t: {"维"} for t in VOC_TYPES},
        known_segments={"新手"},
        items=[],
    )
    assert err and "未知人群" in err


def run_all() -> None:
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    for t in tests:
        t()
        print(f"  ok: {t.__name__}")
    print(f"SEGMENT_TESTS_PASSED ({len(tests)})")


if __name__ == "__main__":
    run_all()

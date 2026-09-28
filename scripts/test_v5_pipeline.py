# -*- coding: utf-8 -*-
"""V5 acceptance tests: rates, dedup, charts, intelligence, hash, excel sheets."""
from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook, load_workbook

from lib.constants import (
    AI_SUMMARY_SHEET_NAME,
    DECISION_SHEET_NAME,
    OVERVIEW_SHEET_NAME,
    RESULT_SHEET_NAME,
    SEGMENT_SHEET_NAME,
    VOC_TYPES,
)
from lib.excel_io import _bar_data_label, write_analysis_workbook
from lib.intelligence_codec import (
    build_intelligence_hash_payload,
    compute_input_hash,
    enrich_intelligence,
    parse_intelligence_output,
    validate_intelligence_payload,
)
from lib.normalize_codec import build_normalize_input_rows, validate_normalize_mappings
from lib.statistics import aggregate_statistics, pick_representative_feedback
from lib.io_util import skill_root


def test_1_percentage_denominator() -> None:
    """analyzed_reviews=80, hits=20 → 25.0% not 20.0%."""
    items = [
        {
            "item_type": "用户满意",
            "dimension": "易清洗",
            "merged_dimension": "易清洗",
            "extracted_summary": f"fb{i}",
            "review_row": i + 1,
        }
        for i in range(20)
    ]
    rows = aggregate_statistics(items, analyzed_reviews=80, consolidate=False)
    assert len(rows) == 1
    assert rows[0]["mention_count"] == 20
    assert rows[0]["mention_rate"] == 25.0


def test_2_review_dedup() -> None:
    """Same review multiple items for same std dim → mention_count=1."""
    items = [
        {
            "item_type": "用户满意",
            "dimension": "易清洗",
            "merged_dimension": "易清洗",
            "extracted_summary": "easy to clean",
            "review_row": 7,
        },
        {
            "item_type": "用户满意",
            "dimension": "易清洗",
            "merged_dimension": "易清洗",
            "extracted_summary": "very easy to clean",
            "review_row": 7,
        },
        {
            "item_type": "用户满意",
            "dimension": "易清洗",
            "merged_dimension": "易清洗",
            "extracted_summary": "cleaning is simple",
            "review_row": 7,
        },
    ]
    rows = aggregate_statistics(items, analyzed_reviews=10, consolidate=False)
    assert rows[0]["mention_count"] == 1
    assert rows[0]["mention_rate"] == 10.0


def test_3_chart_uses_mention_rate() -> None:
    item = {"mention_rate": 25.0, "mention_count": 20}
    assert _bar_data_label(item) == "25%"
    assert "20/80" not in _bar_data_label(item)
    # Hidden chart cell value scale
    assert float(item["mention_rate"]) / 100.0 == 0.25


def test_4_ai_summary_copyable() -> None:
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "in.xlsx"
        out = Path(td) / "out.xlsx"
        wb = Workbook()
        wb.active.title = "Reviews"
        wb["Reviews"].append(["标题", "内容"])
        wb.save(src)
        wb.close()
        text = "AI评论洞察总结\n\n【消费人群】\n· 测试"
        write_analysis_workbook(
            src,
            out,
            summary_rows=[],
            total_reviews=10,
            analyzed_reviews=8,
            overview_summary=text,
            intelligence={"attribute_performance": [], "pain_priorities": [], "opportunities": [], "recommendations": {}},
        )
        wb2 = load_workbook(out)
        assert AI_SUMMARY_SHEET_NAME in wb2.sheetnames
        ws = wb2[AI_SUMMARY_SHEET_NAME]
        assert ws["A1"].value == "AI评论洞察总结"
        assert "消费人群" in str(ws["A2"].value)
        # A2 must not be part of a merged range
        merged = {str(r) for r in ws.merged_cells.ranges}
        assert not any("A2" in m or m.startswith("A2:") for m in merged)
        wb2.close()


def test_5_eight_sections_required() -> None:
    payload = {
        "sections": [{"title": t, "bullets": ["证据不足"]} for t in VOC_TYPES],
        "attribute_performance": [],
        "pain_priorities": [],
        "opportunities": [],
        "recommendations": {
            "priority_improvements": [],
            "keep_strengths": [],
            "explore_opportunities": [],
        },
    }
    assert validate_intelligence_payload(payload) is None
    bad = {
        "sections": [{"title": "消费人群", "bullets": ["x"]}],
        "attribute_performance": [],
        "pain_priorities": [],
        "opportunities": [],
        "recommendations": {
            "priority_improvements": [],
            "keep_strengths": [],
            "explore_opportunities": [],
        },
    }
    err = validate_intelligence_payload(bad)
    assert err and "缺少维" in err


def test_6_unknown_dimension_rejected() -> None:
    summary = [
        {
            "item_type": "未被满足",
            "dimension": "容易断裂",
            "mention_count": 3,
            "mention_rate": 10.0,
        }
    ]
    payload = {
        "sections": [{"title": t, "bullets": ["ok"]} for t in VOC_TYPES],
        "attribute_performance": [],
        "pain_priorities": [
            {
                "dimension": "完全不存在的痛点",
                "severity": "高",
                "priority": "高",
                "reason": "测试",
            }
        ],
        "opportunities": [],
        "recommendations": {
            "priority_improvements": [],
            "keep_strengths": [],
            "explore_opportunities": [],
        },
    }
    enriched, text, err = parse_intelligence_output(
        json.dumps(payload, ensure_ascii=False),
        summary_rows=summary,
        items=[],
        analyzed_reviews=30,
    )
    assert enriched is None and err and "未知" in err


def test_7_opposite_normalize_forbidden() -> None:
    # Prompt hard rule present
    sys_p = (skill_root() / "prompts" / "normalize_system.md").read_text(encoding="utf-8")
    assert "尺寸偏大 ≠ 尺寸偏小" in sys_p or "尺寸偏大≠尺寸偏小" in sys_p.replace(" ", "")
    assert "改进动作一致" in sys_p
    # Codec rejects vague opposite-bucket std names
    err = validate_normalize_mappings(
        [{"类型": "未被满足", "原始维度": "尺寸偏大", "标准维度": "尺寸问题"}],
        {("未被满足", "尺寸偏大")},
    )
    assert err and ("空泛" in err or "相反" in err)


def test_8_step1_workdir_safety() -> None:
    from step1_prepare import _prepare_workdir

    with tempfile.TemporaryDirectory() as td:
        ours = Path(td) / "ours"
        ours.mkdir()
        (ours / "meta.json").write_text(
            json.dumps(
                {
                    "skill_version": "v5",
                    "pipeline": "review_intelligence_v5",
                    "persona_batches": [],
                }
            ),
            encoding="utf-8",
        )
        try:
            _prepare_workdir(ours, force=False)
            raise AssertionError("expected SystemExit without --force")
        except SystemExit as exc:
            assert "非空" in str(exc) or "force" in str(exc).lower()
        assert (ours / "meta.json").exists()
        _prepare_workdir(ours, force=True)
        assert ours.is_dir()
        assert not (ours / "meta.json").exists()


def test_9_summary_hash_invalidation() -> None:
    summary_a = [
        {
            "item_type": "用户满意",
            "dimension": "易清洗",
            "mention_count": 2,
            "mention_rate": 20.0,
            "representative_feedback": "a",
        }
    ]
    summary_b = [
        {
            "item_type": "用户满意",
            "dimension": "易清洗",
            "mention_count": 5,
            "mention_rate": 50.0,
            "representative_feedback": "a",
        }
    ]
    items = [
        {
            "item_type": "用户满意",
            "merged_dimension": "易清洗",
            "signal_type": "满意点",
            "review_row": 1,
        }
    ]
    h1 = compute_input_hash(
        build_intelligence_hash_payload(
            analyzed_reviews=10, summary_rows=summary_a, items=items
        )
    )
    h2 = compute_input_hash(
        build_intelligence_hash_payload(
            analyzed_reviews=10, summary_rows=summary_b, items=items
        )
    )
    h1b = compute_input_hash(
        build_intelligence_hash_payload(
            analyzed_reviews=10, summary_rows=summary_a, items=items
        )
    )
    assert h1 == h1b
    assert h1 != h2


def test_10_excel_sheet_integrity() -> None:
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "in.xlsx"
        out = Path(td) / "out.xlsx"
        wb = Workbook()
        wb.active.title = "Reviews"
        wb.create_sheet("Extra")
        wb["Reviews"].append(["标题", "内容"])
        wb["Reviews"].append(["ok", "good"])
        wb.save(src)
        wb.close()
        rows = [
            {
                "item_type": "消费人群",
                "dimension": "家庭用户",
                "mention_count": 2,
                "mention_rate": 25.0,
                "representative_feedback": "家庭在用",
                "theme_summary": "「家庭用户」提及率 25.0%。",
            }
        ]
        meta = write_analysis_workbook(
            src,
            out,
            summary_rows=rows,
            total_reviews=10,
            analyzed_reviews=8,
            overview_summary="AI评论洞察总结\n\n【消费人群】\n· x",
            intelligence={
                "attribute_performance": [],
                "pain_priorities": [],
                "opportunities": [],
                "recommendations": {
                    "priority_improvements": [],
                    "keep_strengths": [],
                    "explore_opportunities": [],
                },
            },
        )
        names = meta["sheetnames"]
        assert "Reviews" in names and "Extra" in names
        for required in (
            OVERVIEW_SHEET_NAME,
            DECISION_SHEET_NAME,
            SEGMENT_SHEET_NAME,
            AI_SUMMARY_SHEET_NAME,
            RESULT_SHEET_NAME,
        ):
            assert required in names, required
        # analysis sheets at end
        assert names[-5:] == [
            OVERVIEW_SHEET_NAME,
            DECISION_SHEET_NAME,
            SEGMENT_SHEET_NAME,
            AI_SUMMARY_SHEET_NAME,
            RESULT_SHEET_NAME,
        ]
        wb2 = load_workbook(out)
        ov = wb2[OVERVIEW_SHEET_NAME]
        header = str(ov.cell(row=1, column=5).value or "")
        assert "有效分析评论数：8" in header
        assert "30/237" not in header
        wb2.close()


def test_uniform_representative_feedback() -> None:
    items = [
        {"review_row": i, "extracted_summary": f"fb-{i}"}
        for i in range(1, 21)
    ]
    a = pick_representative_feedback(items, max_n=5)
    b = pick_representative_feedback(items, max_n=5)
    assert a == b
    parts = a.split("；")
    assert len(parts) == 5
    assert parts[0] == "fb-1"
    assert parts[-1] == "fb-20"


def test_normalize_multi_refs() -> None:
    items = [
        {
            "item_type": "未被满足",
            "dimension": "安装问题",
            "extracted_summary": f"ref-{i}",
            "review_row": i,
        }
        for i in (1, 2, 3, 4, 5)
    ]
    rows = build_normalize_input_rows(items)
    assert len(rows) == 1
    refs = rows[0]["语义参考"]
    assert isinstance(refs, list)
    assert 1 <= len(refs) <= 3
    assert len(set(refs)) == len(refs)


def test_attribute_rates_dedup() -> None:
    items = [
        {
            "item_type": "用户满意",
            "merged_dimension": "易清洗",
            "review_row": 1,
            "signal_type": "满意点",
        },
        {
            "item_type": "用户满意",
            "merged_dimension": "清洗方便",
            "review_row": 1,
            "signal_type": "满意点",
        },
        {
            "item_type": "未被满足",
            "merged_dimension": "难清洗",
            "review_row": 2,
            "signal_type": "明确问题",
        },
    ]
    summary = aggregate_statistics(
        [
            {**it, "dimension": it["merged_dimension"], "extracted_summary": "x"}
            for it in items
        ],
        analyzed_reviews=10,
    )
    payload = {
        "sections": [{"title": t, "bullets": ["ok"]} for t in VOC_TYPES],
        "attribute_performance": [
            {
                "attribute": "清洁便利性",
                "positive_dimensions": ["易清洗", "清洗方便"],
                "negative_dimensions": ["难清洗"],
                "assessment": "优劣势并存",
            }
        ],
        "pain_priorities": [],
        "opportunities": [],
        "recommendations": {
            "priority_improvements": [],
            "keep_strengths": [],
            "explore_opportunities": [],
        },
    }
    enriched = enrich_intelligence(
        payload, items=items, summary_rows=summary, analyzed_reviews=10
    )
    attr = enriched["attribute_performance"][0]
    # review 1 counted once for positive even if two dims
    assert attr["positive_mention_count"] == 1
    assert attr["positive_mention_rate"] == 10.0
    assert attr["negative_mention_count"] == 1
    assert attr["negative_mention_rate"] == 10.0


def test_signal_required_for_fulfillment() -> None:
    from lib.constants import FULFILLMENT_ALLOWED
    from lib.extract_codec import validate_extract_payload

    payload = {
        "results": [
            {
                "review_id": "R1",
                "items": [
                    {
                        "类型": "用户满意",
                        "单条提炼": "很好洗",
                        "原始维度": "易清洗",
                    }
                ],
            }
        ]
    }
    err = validate_extract_payload(payload, {"R1"}, allowed_types=FULFILLMENT_ALLOWED)
    assert err and "信号类型" in err
    payload["results"][0]["items"][0]["信号类型"] = "满意点"
    assert validate_extract_payload(payload, {"R1"}, allowed_types=FULFILLMENT_ALLOWED) == ""


def run_all() -> None:
    # Also run legacy unit suites where still valid
    import test_excel_export as tex
    import test_label_consolidate as tlc
    import test_normalize_codec as tnc
    import test_summary_codec as tsc

    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    for t in tests:
        t()
        print(f"  ok: {t.__name__}")

    # Updated legacy suites
    tex.test_display_labels()
    print("  ok: tex.test_display_labels")
    tlc.test_long_tail_bucket_only()
    tlc.test_overview_top_n_constant()
    print("  ok: label_consolidate")
    tnc.test_build_rows_dedup_and_count()
    tnc.test_validate_and_apply()
    tnc.test_reject_vague_std()
    print("  ok: normalize_codec")
    tsc.run_all()
    import test_segment_stats as tseg

    tseg.run_all()

    print(f"ALL_V5_TESTS_PASSED core={len(tests)}")


if __name__ == "__main__":
    run_all()

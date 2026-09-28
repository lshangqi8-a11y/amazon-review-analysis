# -*- coding: utf-8 -*-
"""Local assertions for V5 Excel: left charts, right AI summary, decision/AI sheets."""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))

from lib.constants import (
    AI_SUMMARY_SHEET_NAME,
    DECISION_SHEET_NAME,
    OVERVIEW_CHART_TYPES,
    OVERVIEW_SHEET_NAME,
    RESULT_SHEET_NAME,
    SEGMENT_SHEET_NAME,
    VOC_TYPES,
)
from lib.excel_io import PORTRAIT_TYPES, type_display_label, write_analysis_workbook
from lib.statistics import build_theme_insight_summary
from lib.summary_codec import format_summary_text, validate_summary_payload


def _row(item_type, dimension, count, rate, feedback):
    item = {
        "item_type": item_type,
        "dimension": dimension,
        "mention_count": count,
        "mention_rate": rate,
        "representative_feedback": feedback,
    }
    item["theme_summary"] = build_theme_insight_summary(item)
    return item


def _full_summary(*, skip_types=None):
    skip_types = skip_types or set()
    rows = []
    specs = [
        ("消费人群", ["家庭用户", "新手买家", "老年用户", "学生党"], 4),
        ("使用地点", ["厨房", "台面", "庭院", "车内", "室内"], 5),
        ("使用时刻", ["每天", "夜间", "早晨", "碎片时间"], 4),
        ("产品用途", ["日常切割", "厨房备菜", "清洗奶瓶", "娱乐互动"], 4),
        ("使用场景", ["日常清洁养护", "礼赠场合", "出行旅行", "夜间照护"], 4),
        ("购买动机", ["性价比", "品牌信任", "礼品", "朋友推荐"], 4),
        ("用户满意", [f"满意点{i}" for i in range(1, 6)], 5),
        ("未被满足", [f"痛点维度{i}" for i in range(1, 6)], 5),
    ]
    for item_type, dims, _n in specs:
        if item_type in skip_types:
            continue
        for i, dim in enumerate(dims):
            count = max(1, 6 - i) if i < 4 else 1
            rows.append(_row(item_type, dim, count, float(count), f"反馈「{dim}」；补充"))
    return rows


def _sample_summary_text():
    payload = {
        "sections": [
            {"title": t, "bullets": [f"{t}要点一", f"{t}要点二"]} for t in VOC_TYPES
        ]
    }
    assert validate_summary_payload(payload) is None
    return format_summary_text(payload)


def _write_source(path: Path) -> None:
    wb = Workbook()
    wb.active.title = "Reviews"
    wb.create_sheet("Summary")
    ws = wb["Reviews"]
    ws.append(["标题", "内容"])
    ws.append(["good", "works well"])
    wb.save(path)
    wb.close()


def test_dashboard_full() -> None:
    with tempfile.TemporaryDirectory() as td:
        td_path = Path(td)
        src = td_path / "in.xlsx"
        out = td_path / "out.xlsx"
        _write_source(src)
        summary_text = _sample_summary_text()
        meta = write_analysis_workbook(
            src,
            out,
            summary_rows=_full_summary(),
            total_reviews=100,
            analyzed_reviews=100,
            voc_items=40,
            overview_summary=summary_text,
            intelligence={
                "attribute_performance": [],
                "pain_priorities": [],
                "opportunities": [],
                "recommendations": {
                    "priority_improvements": [],
                    "keep_strengths": [],
                    "explore_opportunities": [],
                },
            },
        )
        assert meta["chart_count"] == 4
        assert meta.get("has_overview_summary") is True
        assert str(meta.get("summary_reserve", "")).startswith("O1:U")

        wb = load_workbook(out)
        ov = wb[OVERVIEW_SHEET_NAME]
        assert ov["A1"].value == "Amazon 评论分析总览"
        assert "有效分析评论数：100" in str(ov.cell(row=1, column=5).value)
        assert ov.cell(row=1, column=15).value == "AI总结"
        body = str(ov.cell(row=2, column=15).value or "")
        assert "消费人群" in body and "未被满足" in body
        assert "产品名称" not in body
        assert DECISION_SHEET_NAME in wb.sheetnames
        assert SEGMENT_SHEET_NAME in wb.sheetnames
        assert AI_SUMMARY_SHEET_NAME in wb.sheetnames
        assert RESULT_SHEET_NAME in wb.sheetnames

        joined = []
        for r in range(1, 220):
            for c in range(1, 22):
                v = ov.cell(row=r, column=c).value
                if v is not None:
                    joined.append(str(v))
        text = "\n".join(joined)
        assert "使用地点" in text and "厨房" in text
        assert "评论洞察条目数" not in text
        assert "标准维度数" not in text
        assert "产品名称" not in text
        # Dashboard should not highlight count/total
        assert "（30/237）" not in text
        assert len(ov._charts) == 4
        wb.close()


def test_empty_modules() -> None:
    """Some dimensions have no data: export must not crash and charts shrink."""
    skip = {"使用场景", "购买动机"}
    expected_charts = len([t for t in OVERVIEW_CHART_TYPES if t not in skip])
    with tempfile.TemporaryDirectory() as td:
        td_path = Path(td)
        src = td_path / "in.xlsx"
        out = td_path / "out.xlsx"
        _write_source(src)
        meta = write_analysis_workbook(
            src,
            out,
            summary_rows=_full_summary(skip_types=skip),
            total_reviews=100,
            analyzed_reviews=100,
            voc_items=20,
            overview_summary=_sample_summary_text(),
        )
        assert meta["chart_count"] == expected_charts, meta["chart_count"]

        wb = load_workbook(out)
        ov = wb[OVERVIEW_SHEET_NAME]
        assert len(ov._charts) == expected_charts
        joined = []
        for r in range(1, 220):
            for c in range(1, 22):
                v = ov.cell(row=r, column=c).value
                if v is not None:
                    joined.append(str(v))
        text = "\n".join(joined)
        assert "消费人群" in text and "厨房" in text
        assert "礼赠场合" not in text
        wb.close()


def test_display_labels() -> None:
    assert type_display_label("未被满足") == "未被满足"
    assert PORTRAIT_TYPES == list(OVERVIEW_CHART_TYPES)


if __name__ == "__main__":
    test_display_labels()
    test_dashboard_full()
    test_empty_modules()
    print("ALL_EXCEL_TESTS_PASSED")
